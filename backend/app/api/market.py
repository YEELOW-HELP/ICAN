"""Public regional labour-market API and protected refresh endpoint."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.api.deps import get_privileged_admin as get_current_admin
from app.db.models import AdminUser
from app.db.models_career_kb_mnp import MnpCareer, MnpMarketSnapshot, MnpSalarySnapshot
from app.db.session import get_session
from app.services.market_data.dcz import DATASET_PAGE, SOURCE, refresh_market_data

router = APIRouter(prefix="/v1/mnp", tags=["market-data"])

_EXPORT_COLUMNS = (
    ("Дата зрізу", 14),
    ("Рівень", 15),
    ("Регіон", 30),
    ("Код професії", 23),
    ("Професія", 34),
    ("Категорія", 30),
    ("Кількість вакансій", 20),
    ("Розмір вибірки", 18),
    ("Тренд", 13),
    ("Зарплата P25", 17),
    ("Медіанна зарплата", 21),
    ("Зарплата P75", 17),
    ("Валюта", 11),
    ("Період", 13),
    ("Частка віддаленої роботи", 27),
    ("Доступність entry-level", 25),
    ("Якість даних", 22),
    ("Джерело", 22),
    ("Версія джерела", 35),
)


def _add_export_sheet(workbook: Workbook, title: str, rows: list[tuple]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append([name for name, _width in _EXPORT_COLUMNS])
    header_fill = PatternFill("solid", fgColor="FFC72C")
    thin_line = Side(style="thin", color="E4DED0")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="171713")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin_line)
    sheet.row_dimensions[1].height = 31

    for row in rows:
        snapshot, career, salary = row
        sheet.append([
            snapshot.snapshot_date,
            "Регіон" if snapshot.region else "Україна",
            snapshot.region or "Уся Україна",
            career.code,
            career.canonical_name_uk,
            career.career_family.name_uk if career.career_family else None,
            snapshot.vacancy_count,
            snapshot.sample_size,
            snapshot.demand_trend,
            salary.percentile_25 if salary else None,
            salary.median if salary else None,
            salary.percentile_75 if salary else None,
            salary.currency if salary else None,
            salary.period if salary else None,
            snapshot.remote_share,
            snapshot.entry_level_availability,
            snapshot.data_quality,
            snapshot.source,
            snapshot.source_version,
        ])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for index, (_name, width) in enumerate(_EXPORT_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = "DD.MM.YYYY"
        for index in (6, 7, 9, 10, 11):
            row[index].number_format = '#,##0.00'
        for index in (14, 15):
            row[index].number_format = "0.00%"
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="FAF8F2")


@router.get("/market/export.xlsx")
async def export_market_data(session: AsyncSession = Depends(get_session)):
    """Export every persisted row from the latest official market snapshot."""

    latest = await session.scalar(
        select(func.max(MnpMarketSnapshot.snapshot_date)).where(MnpMarketSnapshot.source == SOURCE)
    )
    if latest is None:
        raise HTTPException(status_code=404, detail="Ринкові дані ще не синхронізовано")

    rows = (await session.execute(
        select(MnpMarketSnapshot, MnpCareer, MnpSalarySnapshot)
        .options(selectinload(MnpCareer.career_family))
        .join(MnpCareer, MnpCareer.id == MnpMarketSnapshot.career_id)
        .outerjoin(MnpSalarySnapshot, MnpSalarySnapshot.market_snapshot_id == MnpMarketSnapshot.id)
        .where(
            MnpMarketSnapshot.source == SOURCE,
            MnpMarketSnapshot.snapshot_date == latest,
        )
        .order_by(
            MnpMarketSnapshot.region.asc().nullsfirst(),
            MnpMarketSnapshot.vacancy_count.desc(),
            MnpCareer.canonical_name_uk,
        )
    )).all()

    workbook = Workbook()
    workbook.remove(workbook.active)
    all_rows = list(rows)
    _add_export_sheet(workbook, "Усі дані", all_rows)
    _add_export_sheet(workbook, "Україна", [row for row in all_rows if row[0].region is None])
    _add_export_sheet(workbook, "Регіони", [row for row in all_rows if row[0].region is not None])

    about = workbook.create_sheet("Про експорт")
    about.sheet_view.showGridLines = False
    details = (
        ("Експортовано", datetime.now(timezone.utc).strftime("%d.%m.%Y %H:%M UTC")),
        ("Дата зрізу", latest.strftime("%d.%m.%Y")),
        ("Кількість записів", len(all_rows)),
        ("Джерело", "Державна служба зайнятості України / Data.gov.ua"),
        ("Посилання", DATASET_PAGE),
        ("Методологія", "Експорт містить усі збережені агреговані записи останнього зрізу: професія × регіон. Назви вакансій точно зіставлені з каталогом професій ICAN."),
        ("Обмеження", "Первинні тексти та контакти окремих вакансій поточна модель не зберігає."),
    )
    for key, value in details:
        about.append([key, value])
    about.column_dimensions["A"].width = 23
    about.column_dimensions["B"].width = 105
    for cell in about[1]:
        cell.fill = PatternFill("solid", fgColor="FFC72C")
    for row in about.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"ican-market-data-{latest.isoformat()}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/market/overview")
async def market_overview(
    region: str | None = Query(default=None, max_length=128),
    limit: int = Query(default=12, ge=1, le=30),
    session: AsyncSession = Depends(get_session),
):
    latest = await session.scalar(
        select(func.max(MnpMarketSnapshot.snapshot_date)).where(MnpMarketSnapshot.source == SOURCE)
    )
    if latest is None:
        return {
            "available": False,
            "message": "Ринкові дані ще синхронізуються",
            "source": "Державна служба зайнятості України",
            "source_url": DATASET_PAGE,
            "regions": [],
            "region_summaries": [],
            "items": [],
        }

    regions = list((await session.scalars(
        select(MnpMarketSnapshot.region)
        .where(
            MnpMarketSnapshot.source == SOURCE,
            MnpMarketSnapshot.snapshot_date == latest,
            MnpMarketSnapshot.region.is_not(None),
        )
        .distinct()
        .order_by(MnpMarketSnapshot.region)
    )).all())
    selected_region = region if region in regions else None
    region_filter = (
        MnpMarketSnapshot.region == selected_region
        if selected_region else MnpMarketSnapshot.region.is_(None)
    )
    rows = (await session.execute(
        select(MnpMarketSnapshot, MnpCareer, MnpSalarySnapshot)
        .options(selectinload(MnpCareer.career_family))
        .join(MnpCareer, MnpCareer.id == MnpMarketSnapshot.career_id)
        .outerjoin(MnpSalarySnapshot, MnpSalarySnapshot.market_snapshot_id == MnpMarketSnapshot.id)
        .where(
            MnpMarketSnapshot.source == SOURCE,
            MnpMarketSnapshot.snapshot_date == latest,
            region_filter,
        )
        .order_by(MnpMarketSnapshot.vacancy_count.desc(), MnpCareer.canonical_name_uk)
        .limit(limit)
    )).all()
    total = int((await session.scalar(
        select(func.sum(MnpMarketSnapshot.vacancy_count)).where(
            MnpMarketSnapshot.source == SOURCE,
            MnpMarketSnapshot.snapshot_date == latest,
            region_filter,
        )
    )) or 0)
    region_summary_rows = (await session.execute(
        select(
            MnpMarketSnapshot.region,
            func.sum(MnpMarketSnapshot.vacancy_count),
            func.avg(MnpSalarySnapshot.median),
        )
        .outerjoin(MnpSalarySnapshot, MnpSalarySnapshot.market_snapshot_id == MnpMarketSnapshot.id)
        .where(
            MnpMarketSnapshot.source == SOURCE,
            MnpMarketSnapshot.snapshot_date == latest,
            MnpMarketSnapshot.region.is_not(None),
        )
        .group_by(MnpMarketSnapshot.region)
        .order_by(MnpMarketSnapshot.region)
    )).all()
    return {
        "available": True,
        "snapshot_date": latest.isoformat(),
        "selected_region": selected_region,
        "total_vacancies": total,
        "source": "Державна служба зайнятості України",
        "source_url": DATASET_PAGE,
        "methodology_note": "Показано вакансії, назви яких точно зіставлені з перевіреним каталогом професій ICAN.",
        "regions": regions,
        "region_summaries": [
            {
                "region": region_name,
                "vacancy_count": int(vacancy_count or 0),
                "salary_average": round(float(salary_average)) if salary_average is not None else None,
            }
            for region_name, vacancy_count, salary_average in region_summary_rows
        ],
        "items": [
            {
                "career_id": str(career.id),
                "career_code": career.code,
                "name": career.canonical_name_uk,
                "family": career.career_family.name_uk if career.career_family else None,
                "vacancy_count": snapshot.vacancy_count or 0,
                "trend": snapshot.demand_trend,
                "salary_median": salary.median if salary else None,
                "salary_p25": salary.percentile_25 if salary else None,
                "salary_p75": salary.percentile_75 if salary else None,
            }
            for snapshot, career, salary in rows
        ],
    }


@router.post("/admin/market/refresh")
async def refresh_market(
    _admin: AdminUser = Depends(get_current_admin),
    session: AsyncSession = Depends(get_session),
):
    return await refresh_market_data(session)
