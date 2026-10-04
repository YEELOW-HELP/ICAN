"""Administrator-only weekly Excel reports for the shared client workspace."""

from datetime import datetime, timezone
from io import BytesIO

import httpx
import pytest
from openpyxl import load_workbook

from app.mongo_runtime.core import ADMIN, MANAGER, current_staff, database
from app.mongo_runtime.main import app
from tests.test_mongo_dropbox_cv import Collection, Database


def _section_row(sheet, title: str) -> int:
    return next(row for row in range(1, sheet.max_row + 1) if str(sheet.cell(row, 1).value).startswith(title))


@pytest.mark.asyncio
async def test_admin_downloads_weekly_report_and_manager_is_denied():
    db = Database()
    db.admin_users = Collection([
        {"_id": 1, "email": "admin@example.com", "full_name": "Адміністратор",
         "role": ADMIN, "is_active": True},
        {"_id": 7, "email": "manager@example.com", "full_name": "Консультант",
         "role": MANAGER, "is_active": True},
    ])
    db.mnp_client_request_types = Collection([
        {"_id": "request-1", "name": "Пошук роботи", "is_active": True},
    ])
    db.mnp_employment_stages = Collection([
        {"_id": "result-1", "name": "Проходить співбесіду", "is_active": True},
    ])
    db.mnp_persons = Collection([
        {
            "_id": "person-1", "first_name": "Олена", "last_name": "Тестова",
            "phone": "+380501234567", "email": "olena@example.com", "city": "Харків",
            "status": "case", "workflow_stage": "needs_contact", "needs_contact": True,
            "responsible_staff_id": 7, "client_request_ids": ["request-1"],
            "work_format": "remote", "employment_type": "part_time",
            "employment_stage_id": "result-1", "employment_offer_text": "Добір вакансій",
            "referral_source": "instagram", "tags": [{"name": "Excel"}, {"name": "CRM"}],
            "created_at": "2026-10-01T08:00:00Z", "updated_at": "2026-10-02T09:00:00Z",
        },
        {
            "_id": "person-2", "first_name": "Інший", "status": "active",
            "workflow_stage": "in_progress", "needs_contact": False,
            "created_at": "2026-09-01T08:00:00Z", "updated_at": "2026-09-10T08:00:00Z",
        },
    ])
    db.mnp_client_interactions = Collection([
        {
            "_id": "event-1", "person_id": "person-1", "staff_id": 7,
            "staff_role": MANAGER, "action": "workflow_updated",
            "action_uk": "Оновлено супровід або статус", "details": {},
            "occurred_at": datetime(2026, 10, 1, 9, 0, tzinfo=timezone.utc),
        },
        {
            "_id": "event-2", "person_id": "person-1", "staff_id": 7,
            "staff_role": MANAGER, "action": "fact_updated",
            "action_uk": "Оновлено дані профілю", "details": {"fact_type": "skills"},
            "occurred_at": datetime(2026, 10, 2, 10, 0, tzinfo=timezone.utc),
        },
    ])
    app.dependency_overrides[database] = lambda: db
    app.dependency_overrides[current_staff] = lambda: {
        "_id": 1, "email": "admin@example.com", "role": ADMIN, "is_active": True,
    }
    try:
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
            response = await client.get(
                "/v1/mnp/admin/persons/report.xlsx", params={"week_start": "2026-09-28"},
            )
            assert response.status_code == 200, response.text
            assert response.headers["content-type"].startswith(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            assert "yellow-hub-weekly-2026-09-28" in response.headers["content-disposition"]
            workbook = load_workbook(BytesIO(response.content))
            sheet = workbook["Тижневий звіт"]
            assert sheet["A1"].value == "Yellow Hub · Тижневий звіт"
            assert "Опрацьовано клієнтів: 1" in sheet["A2"].value
            assert "Дій: 2" in sheet["A2"].value
            assert "Усього клієнтів: 2" in sheet["A2"].value

            team_row = _section_row(sheet, "РОБОТА КОМАНДИ ЗА ТИЖДЕНЬ")
            team_headers = [cell.value for cell in sheet[team_row + 1]][:5]
            team_values = [cell.value for cell in sheet[team_row + 2]][:5]
            team = dict(zip(team_headers, team_values))
            assert team["Працівник"] == "Консультант"
            assert team["Опрацьовано клієнтів"] == 1
            assert team["Кількість дій"] == 2

            weekly_row = _section_row(sheet, "КЛІЄНТИ, ОПРАЦЬОВАНІ ЗА ТИЖДЕНЬ")
            weekly_headers = [cell.value for cell in sheet[weekly_row + 1]][:9]
            weekly_values = [cell.value for cell in sheet[weekly_row + 2]][:9]
            weekly = dict(zip(weekly_headers, weekly_values))
            assert weekly["Клієнт"] == "Олена Тестова"
            assert weekly["Остання взаємодія"] == "02.10.2026 10:00"
            assert weekly["Хто працював"] == "Консультант"
            assert weekly["Кількість дій"] == 2
            assert "навички" in weekly["Виконана робота"]

            all_row = _section_row(sheet, "РЕШТА КЛІЄНТІВ")
            all_headers = [cell.value for cell in sheet[all_row + 1]][:23]
            first_client = dict(zip(all_headers, [cell.value for cell in sheet[all_row + 2]][:23]))
            assert "ID клієнта" not in all_headers
            assert first_client["Клієнт"] == "Інший"
            assert first_client["Остання взаємодія"] == "10.09.2026 08:00"
            assert sheet.max_row == all_row + 2

            app.dependency_overrides[current_staff] = lambda: {
                "_id": 7, "email": "manager@example.com", "role": MANAGER, "is_active": True,
            }
            forbidden = await client.get("/v1/mnp/admin/persons/report.xlsx")
            assert forbidden.status_code == 403
    finally:
        app.dependency_overrides.pop(database, None)
        app.dependency_overrides.pop(current_staff, None)
