from datetime import date
from io import BytesIO

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from app.api.main import app
from app.db.models_career_kb_mnp import (
    CareerLifecycleStatus,
    MnpCareer,
    MnpCareerFamily,
    MnpMarketSnapshot,
    MnpSalarySnapshot,
)
from app.db.session import get_session
from app.services.market_data.dcz import iter_vacancies, normalize_region, normalize_title, parse_salary


async def _chunks(payload: bytes):
    for index in range(0, len(payload), 17):
        yield payload[index:index + 17]


@pytest.mark.asyncio
async def test_dcz_xml_is_streamed_and_normalized():
    xml = b"""<?xml version='1.0' encoding='utf-8'?><jobs>
      <job id='1'><name>![CDATA[Content manager]]</name>
      <region>![CDATA[Kharkiv region, Kharkiv district]]</region>
      <salary>![CDATA[25 000 UAH per month]]</salary></job>
      <job id='2'><name>Accountant</name><region>Kyiv</region><salary>not specified</salary></job>
    </jobs>"""
    # ASCII fixture exercises the source's literal ![CDATA[...]] wrapper;
    # Ukrainian normalization is covered separately below.
    rows = [row async for row in iter_vacancies(_chunks(xml))]
    assert len(rows) == 2
    assert rows[0].title == "Content manager"
    assert rows[0].salary == 25_000
    assert rows[1].salary is None
    assert normalize_title("  SMM-менеджер ") == "smm менеджер"
    assert normalize_region("Харківська область, Харківський район") == "Харківська область"
    assert parse_salary("![CDATA[18 500₴ за місяць]]") == 18_500


@pytest.mark.asyncio
async def test_market_overview_returns_latest_regional_snapshot(session_factory):
    async with session_factory() as session:
        family = MnpCareerFamily(code="marketing", name_uk="Маркетинг", name_en="Marketing")
        session.add(family)
        await session.flush()
        career = MnpCareer(
            code="content_manager",
            canonical_name_uk="Контент-менеджер",
            canonical_name_en="Content Manager",
            description_short_uk="Створює та керує контентом",
            career_family_id=family.id,
            status=CareerLifecycleStatus.ACTIVE,
        )
        session.add(career)
        await session.flush()
        snapshot = MnpMarketSnapshot(
            career_id=career.id,
            country="UA",
            region="Харківська область",
            snapshot_date=date(2026, 9, 4),
            source="data.gov.ua/dcz",
            source_version="0e7841b3bf854c98b70f46fe3dd031e9",
            data_quality="OFFICIAL_OPEN_DATA",
            sample_size=17,
            vacancy_count=17,
            demand_trend="up",
        )
        session.add(snapshot)
        await session.flush()
        session.add(MnpSalarySnapshot(
            market_snapshot_id=snapshot.id,
            currency="UAH", period="month", percentile_25=18_000,
            median=25_000, percentile_75=32_000,
        ))
        await session.commit()

    async def override_get_session():
        async with session_factory() as session:
            yield session

    app.dependency_overrides[get_session] = override_get_session
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            response = await client.get(
                "/v1/mnp/market/overview",
                params={"region": "Харківська область"},
            )
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 200
    body = response.json()
    assert body["available"] is True
    assert body["snapshot_date"] == "2026-09-04"
    assert body["total_vacancies"] == 17
    assert body["region_summaries"] == [{
        "region": "Харківська область",
        "vacancy_count": 17,
        "salary_average": 25_000,
    }]
    assert body["items"][0]["name"] == "Контент-менеджер"
    assert body["items"][0]["salary_median"] == 25_000


@pytest.mark.asyncio
async def test_market_export_contains_every_latest_snapshot_row(session_factory):
    async with session_factory() as session:
        family = MnpCareerFamily(code="finance", name_uk="Фінанси", name_en="Finance")
        session.add(family)
        await session.flush()
        career = MnpCareer(
            code="accountant",
            canonical_name_uk="Бухгалтер",
            canonical_name_en="Accountant",
            description_short_uk="Веде облік",
            career_family_id=family.id,
            status=CareerLifecycleStatus.ACTIVE,
        )
        session.add(career)
        await session.flush()
        for region, count in ((None, 11), ("Київська область", 7)):
            snapshot = MnpMarketSnapshot(
                career_id=career.id,
                country="UA", region=region, snapshot_date=date(2026, 9, 5),
                source="data.gov.ua/dcz", source_version="snapshot-v1",
                data_quality="OFFICIAL_OPEN_DATA", sample_size=count,
                vacancy_count=count, demand_trend="flat",
            )
            session.add(snapshot)
            await session.flush()
            session.add(MnpSalarySnapshot(
                market_snapshot_id=snapshot.id, currency="UAH", period="month",
                percentile_25=15_000, median=20_000, percentile_75=25_000,
            ))
        await session.commit()

    async def override_get_session():
        async with session_factory() as session:
            yield session

    app.dependency_overrides[get_session] = override_get_session
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            response = await client.get("/v1/mnp/market/export.xlsx")
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 200
    assert response.content.startswith(b"PK")
    assert "ican-market-data-2026-09-05.xlsx" in response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    assert workbook.sheetnames == ["Усі дані", "Україна", "Регіони", "Про експорт"]
    rows = list(workbook["Усі дані"].iter_rows(values_only=True))
    assert len(rows) == 3
    assert rows[0][0:7] == (
        "Дата зрізу", "Рівень", "Регіон", "Код професії", "Професія", "Категорія", "Кількість вакансій",
    )
    assert {row[2] for row in rows[1:]} == {"Уся Україна", "Київська область"}
    assert {row[6] for row in rows[1:]} == {11, 7}
